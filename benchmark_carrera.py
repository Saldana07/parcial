"""
Repite la Carrera sin ventana para cada tamaño, nº de hilos/procesos y generaciones, y guarda
resultados/ejecuciones.csv y resultados/comparacion_resultados.xlsx (tablas y gráficas).
Uso: python benchmark_carrera.py [--tamanos 100,500,1000 --trabajadores 4,8,16 --generaciones 200,500 --repeticiones 5]
"""

import argparse
import csv
import os
import statistics
import time
from types import SimpleNamespace

import config
from core import grid as grid_mod
from gui.race import Race

VERSIONES = [1, 2, 3, 4]
NOMBRES = {1: "V1 Secuencial", 2: "V2 Hilos", 3: "V3 Procesos", 4: "V4 NumPy"}
CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resultados")
TIMEOUT_S = 1800         # si una carrera tarda más que esto (30 min), se cancela

# Colores de las gráficas: los mismos que usa la GUI para cada versión (hex RGB).
COLORES = {1: "96AABE", 2: "5694FF", 3: "FF9942", 4: "5CCE78"}
COLOR_BASE = "7F7F7F"                                  # línea de referencia "1×" (= igual que V1)
COL_TIEMPO = {1: "D", 2: "E", 3: "F", 4: "G"}          # columna de cada versión en las hojas de tamaño


def lista_enteros(texto):
    return [int(x) for x in texto.split(",") if x.strip()]


def configuraciones(filas):
    """Combinaciones (tamaño, hilos/proc, generaciones) en el orden en que se midieron."""
    vistas = []
    for f in filas:
        k = (f["tamano"], f["trabajadores"], f["generaciones"])
        if k not in vistas:
            vistas.append(k)
    return vistas


# --- 1. medir
def una_carrera(tablero, trabajadores, generaciones):
    """Corre UNA carrera (igual que la GUI). Devuelve {versión: segundos} y {versión: ¿igual a V1?}."""
    carrera = Race(VERSIONES, tablero, trabajadores, generaciones)
    inicio = time.perf_counter()
    try:
        while not carrera.finished:
            carrera.poll()
            if time.perf_counter() - inicio > TIMEOUT_S:
                raise TimeoutError(f"la carrera superó {TIMEOUT_S} s")
            time.sleep(0.005)
        carrera.poll()
        if carrera.errors:
            raise RuntimeError(f"fallaron versiones: {carrera.errors}")
        tiempos = {v: carrera.results[v][0] for v in VERSIONES}
        base = carrera.results[1][1]
        correcto = {v: grid_mod.equals(base, carrera.results[v][1]) for v in VERSIONES}
        return tiempos, correcto
    finally:
        carrera.cancel()        # limpia los procesos hijos


def medir(filas, tamanos, trabajadores, generaciones, repeticiones, semilla, ruta_csv):
    """Agrega a `filas` una entrada por carrera. Guarda el CSV tras cada una (por si se interrumpe)."""
    total = len(tamanos) * len(trabajadores) * len(generaciones) * repeticiones
    hecha = 0
    for tamano in tamanos:
        tablero = grid_mod.new_grid(tamano, tamano)
        grid_mod.randomize(tablero, semilla, config.DENSIDAD)     # el mismo tablero para cada tamaño
        for p in trabajadores:
            for g in generaciones:
                for rep in range(1, repeticiones + 1):
                    hecha += 1
                    print(f"[{hecha}/{total}] matriz {tamano}x{tamano}  hilos/proc={p}  generaciones={g}  "
                          f"ejecución {rep}/{repeticiones} ...", end=" ", flush=True)
                    tiempos, correcto = una_carrera(tablero, p, g)
                    print("  ".join(f"V{v}={tiempos[v]:.2f}s" for v in VERSIONES))
                    filas.append({"tamano": tamano, "trabajadores": p, "generaciones": g, "ejecucion": rep,
                                  "tiempos": tiempos, "correcto": all(correcto.values())})
                    guardar_csv(filas, ruta_csv)


# --- 2. guardar
def guardar_csv(filas, ruta):
    with open(ruta, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["matriz", "hilos_proc", "generaciones", "ejecucion",
                    "V1_s", "V2_s", "V3_s", "V4_s", "correcto"])
        for fila in filas:
            w.writerow([f"{fila['tamano']}x{fila['tamano']}", fila["trabajadores"], fila["generaciones"],
                        fila["ejecucion"], *[f"{fila['tiempos'][v]:.4f}" for v in VERSIONES], fila["correcto"]])


def leer_csv(ruta):
    """Lee resultados/ejecuciones.csv y devuelve la misma lista de `filas` que produce medir()."""
    filas = []
    with open(ruta, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            filas.append({"tamano": int(r["matriz"].split("x")[0]), "trabajadores": int(r["hilos_proc"]),
                          "generaciones": int(r["generaciones"]), "ejecucion": int(r["ejecucion"]),
                          "tiempos": {v: float(r[f"V{v}_s"]) for v in VERSIONES},
                          "correcto": r["correcto"].strip().lower() == "true"})
    return filas


# --- gráficas de Excel
def _estilo_ejes(chart, titulo, x_titulo, y_titulo, ancho_cm, alto_cm, log=False, formato_y=None):
    """Título, ejes, leyenda abajo y tamaño de una gráfica. `log` = escala logarítmica en Y."""
    chart.title = titulo
    chart.x_axis.title = x_titulo
    chart.y_axis.title = y_titulo
    chart.width, chart.height = ancho_cm, alto_cm
    chart.x_axis.delete = False             # sin esto, algunas versiones de Excel ocultan los ejes
    chart.y_axis.delete = False
    chart.legend.position = "b"
    if log:
        chart.y_axis.scaling.logBase = 10   # V4 es ~50x más rápida que V1: en escala lineal no se vería
    if formato_y:
        chart.y_axis.number_format = formato_y


def _pintar_barras(chart, colores):
    from openpyxl.chart.shapes import GraphicalProperties
    for serie, color in zip(chart.series, colores):
        serie.graphicalProperties = GraphicalProperties(solidFill=color)
        serie.graphicalProperties.line.solidFill = color


def _pintar_lineas(chart, colores, punteada_la_ultima=False, primera_gruesa=False):
    """Línea con marcador del color de cada versión; la última (referencia 1×) va gris y punteada."""
    from openpyxl.chart.marker import Marker
    from openpyxl.chart.shapes import GraphicalProperties
    for i, (serie, color) in enumerate(zip(chart.series, colores)):
        es_base = punteada_la_ultima and i == len(colores) - 1
        serie.graphicalProperties = GraphicalProperties()
        serie.graphicalProperties.line.solidFill = color
        serie.graphicalProperties.line.width = 22000 if es_base else (60000 if primera_gruesa and i == 0 else 28000)
        serie.smooth = False
        if es_base:
            serie.graphicalProperties.line.dashStyle = "dash"
            serie.marker = Marker(symbol="none")
        else:
            serie.marker = Marker(symbol="circle", size=7)
            serie.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            serie.marker.graphicalProperties.line.solidFill = color


def _graficas_tamano(ws, n_configs):
    """Dos gráficas a la derecha de las tablas de una hoja de tamaño, leyendo la TABLA DE PROMEDIOS:"""
    from openpyxl.chart import BarChart, LineChart, Reference

    fin = 2 + n_configs                                   # última fila de la tabla de promedios
    etiquetas = Reference(ws, min_col=12, min_row=3, max_row=fin)

    tiempos = BarChart()
    tiempos.type, tiempos.grouping = "col", "clustered"
    tiempos.add_data(Reference(ws, min_col=4, max_col=7, min_row=2, max_row=fin), titles_from_data=True)
    tiempos.set_categories(etiquetas)
    _estilo_ejes(tiempos, "Tiempo promedio por configuración (escala logarítmica)",
                 "Hilos/procesos · generaciones", "Tiempo (s, escala log)", 22, 9.5, log=True)
    _pintar_barras(tiempos, [COLORES[v] for v in VERSIONES])
    ws.add_chart(tiempos, "O2")

    speedup = BarChart()
    speedup.type, speedup.grouping = "col", "clustered"
    speedup.add_data(Reference(ws, min_col=8, max_col=10, min_row=2, max_row=fin), titles_from_data=True)
    speedup.set_categories(etiquetas)
    _estilo_ejes(speedup, "Speedup contra V1 (escala logarítmica)",
                 "Hilos/procesos · generaciones", "Speedup (× contra V1, escala log)", 22, 9.5, log=True,
                 formato_y='0.0"×"')
    _pintar_barras(speedup, [COLORES[v] for v in (2, 3, 4)])
    base = LineChart()
    base.add_data(Reference(ws, min_col=13, min_row=2, max_row=fin), titles_from_data=True)
    base.set_categories(etiquetas)
    _pintar_lineas(base, [COLOR_BASE], punteada_la_ultima=True)
    speedup += base                                       # barras + línea de referencia 1× en el mismo gráfico
    ws.add_chart(speedup, "O22")


def _hoja_resumen(wb, filas, hojas, E):
    """Hoja "Resumen" con las gráficas del enunciado, que leen las celdas de las tablas."""
    from openpyxl.chart import LineChart, Reference

    ws = wb.create_sheet("Resumen")
    gmax = max(f["generaciones"] for f in filas)
    tamanos = sorted(hojas)
    trabajadores = sorted({f["trabajadores"] for f in filas})
    existen = {(f["tamano"], f["trabajadores"], f["generaciones"]) for f in filas}

    def encabezado_col(fila, col, texto):
        c = ws.cell(row=fila, column=col, value=texto)
        c.font, c.fill, c.alignment, c.border = E.hdr_font, E.hdr_fill, E.centro, E.borde

    def cabecera_bloque(fila, etiqueta1, valor1, valor2):
        """Fila con los parámetros del bloque: etiqueta1 | valor1 | Generaciones | valor2 (azul = dato)."""
        for col, texto in ((1, etiqueta1), (3, "Generaciones:")):
            ws.cell(row=fila, column=col, value=texto).font = E.negrita
        for col, valor in ((2, valor1), (4, valor2)):
            c = ws.cell(row=fila, column=col, value=valor)
            E.estilo(c, E.dato)

    def tabla(fila, titulos, filas_tabla, formatos):
        for i, t in enumerate(titulos, 1):
            encabezado_col(fila, i, t)
        ws.row_dimensions[fila].height = 32
        for r, valores in enumerate(filas_tabla, fila + 1):
            for c, (valor, formato) in enumerate(zip(valores, formatos), 1):
                celda = ws.cell(row=r, column=c, value=valor)
                es_dato = not (isinstance(valor, str) and valor.startswith("="))
                E.estilo(celda, E.dato if es_dato else E.calc, formato)

    def seccion(fila, texto):
        ws.cell(row=fila, column=1, value=texto).font = E.titulo
        return fila + 1

    ws.cell(row=1, column=1, value="Resumen — gráficas de rendimiento").font = E.titulo
    ws.cell(row=2, column=1, value=(
        f"Todas las tablas promedian las ejecuciones de las hojas de tamaño con {gmax} generaciones (el mayor valor "
        "medido). Las celdas en azul son datos: puedes cambiarlas por otro valor que exista en las hojas de tamaño "
        "y las gráficas se actualizan.")).font = E.calc
    ws.cell(row=3, column=1, value=(
        "Escala logarítmica en todas las gráficas: V4 es ~50 veces más rápida que V1 y en escala lineal las demás "
        "no se distinguirían. La línea punteada del speedup marca 1× (igual de rápida que V1).")).font = E.calc
    fila = 5

    # --- sección 1: tiempo vs tamaño de la matriz, una gráfica por cantidad de hilos/procesos
    fila = seccion(fila, "1. Tiempo vs tamaño de la matriz")
    for p in trabajadores:
        tams = [t for t in tamanos if (t, p, gmax) in existen]
        if not tams:
            continue
        cabecera_bloque(fila, "Hilos/proc.:", p, gmax)
        datos = []
        for t in tams:
            hoja, a, b = hojas[t]
            fila_datos = []
            for v in VERSIONES:
                col = COL_TIEMPO[v]
                fila_datos.append(
                    f"=AVERAGEIFS('{hoja}'!${col}${a}:${col}${b},'{hoja}'!$A${a}:$A${b},$B${fila},"
                    f"'{hoja}'!$B${a}:$B${b},$D${fila})")
            datos.append([f"{t} × {t}", *fila_datos])
        tabla(fila + 1, ["Matriz", *[NOMBRES[v] for v in VERSIONES]], datos,
              [None] + ['0.000" s"'] * 4)
        g = LineChart()
        g.add_data(Reference(ws, min_col=2, max_col=5, min_row=fila + 1, max_row=fila + 1 + len(tams)),
                   titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=fila + 2, max_row=fila + 1 + len(tams)))
        _estilo_ejes(g, "Tiempo vs tamaño de la matriz (escala logarítmica)", "Tamaño de la matriz",
                     "Tiempo (s, escala log)", 15, 7.5, log=True)
        _pintar_lineas(g, [COLORES[v] for v in VERSIONES], primera_gruesa=True)
        ws.add_chart(g, f"G{fila}")
        fila += max(len(tams) + 3, 17)
    fila += 1

    # --- sección 2: tiempo y speedup vs hilos/procesos, una pareja de gráficas por tamaño
    fila = seccion(fila, "2. Tiempo y speedup vs número de hilos/procesos")
    for t in tamanos:
        ps = [p for p in trabajadores if (t, p, gmax) in existen]
        if not ps:
            continue
        hoja, a, b = hojas[t]
        # --- tiempo
        cabecera_bloque(fila, "Matriz:", f"{t} × {t}", gmax)
        datos = []
        for i, p in enumerate(ps):
            r = fila + 2 + i
            fila_datos = []
            for v in VERSIONES:
                col = COL_TIEMPO[v]
                fila_datos.append(
                    f"=AVERAGEIFS('{hoja}'!${col}${a}:${col}${b},'{hoja}'!$A${a}:$A${b},$A{r},"
                    f"'{hoja}'!$B${a}:$B${b},$D${fila})")
            datos.append([p, *fila_datos])
        tabla(fila + 1, ["Hilos/proc.", *[NOMBRES[v] for v in VERSIONES]], datos, [None] + ['0.000" s"'] * 4)
        g = LineChart()
        g.add_data(Reference(ws, min_col=2, max_col=5, min_row=fila + 1, max_row=fila + 1 + len(ps)),
                   titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=fila + 2, max_row=fila + 1 + len(ps)))
        _estilo_ejes(g, f"Tiempo vs hilos/procesos — matriz {t} × {t} (escala logarítmica)",
                     "Hilos / procesos", "Tiempo (s, escala log)", 15, 7.5, log=True)
        _pintar_lineas(g, [COLORES[v] for v in VERSIONES], primera_gruesa=True)
        ws.add_chart(g, f"G{fila}")
        fila_tiempo = fila
        fila += max(len(ps) + 3, 17)

        # --- speedup (fórmulas sobre la tabla de tiempos de arriba) + línea de referencia 1×
        cabecera_bloque(fila, "Matriz:", f"{t} × {t}", gmax)
        datos = []
        for i, p in enumerate(ps):
            rt = fila_tiempo + 2 + i                      # fila de esta configuración en la tabla de tiempos
            datos.append([f"=A{rt}", f"=$B{rt}/C{rt}", f"=$B{rt}/D{rt}", f"=$B{rt}/E{rt}", 1])
        tabla(fila + 1, ["Hilos/proc.", "Speedup V2", "Speedup V3", "Speedup V4", "Base V1 (1×)"], datos,
              [None] + ['0.00"×"'] * 4)
        g = LineChart()
        g.add_data(Reference(ws, min_col=2, max_col=5, min_row=fila + 1, max_row=fila + 1 + len(ps)),
                   titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=fila + 2, max_row=fila + 1 + len(ps)))
        _estilo_ejes(g, f"Speedup vs hilos/procesos — matriz {t} × {t} (escala logarítmica)", "Hilos / procesos",
                     "Speedup (× contra V1, escala log)", 15, 7.5, log=True, formato_y='0.0"×"')
        _pintar_lineas(g, [COLORES[2], COLORES[3], COLORES[4], COLOR_BASE], punteada_la_ultima=True)
        ws.add_chart(g, f"G{fila}")
        fila += max(len(ps) + 3, 17) + 1

    ws.cell(row=fila, column=1, value="Datos medidos en azul; fórmulas en negro. La columna 'Base V1 (1×)' es un valor fijo "
                                      "(1) que solo sirve de línea de referencia en la gráfica de speedup.").font = E.calc
    for i, ancho in enumerate([14, 15, 15, 15, 15, 3], 1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    return ws


def guardar_xlsx(filas, ruta):
    """Un libro con UNA HOJA POR TAMAÑO de matriz ("100x100", "500x500", ...) y una hoja "Resumen"."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    hdr_font = Font(name=F, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3A5F")
    titulo = Font(name=F, bold=True, size=12)
    dato = Font(name=F, color="0000FF")          # azul = dato medido
    calc = Font(name=F)                          # negro = fórmula
    banda = PatternFill("solid", fgColor="EEF3F8")
    lado = Side(style="thin", color="BFBFBF")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def encabezado(ws, fila, etiquetas):
        for i, texto in enumerate(etiquetas, 1):
            c = ws.cell(row=fila, column=i, value=texto)
            c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, centro, borde
        ws.row_dimensions[fila].height = 32

    def estilo(celda, fuente, formato=None):
        celda.font, celda.alignment, celda.border = fuente, centro, borde
        if formato:
            celda.number_format = formato

    def pagina_horizontal(ws):
        """Para imprimir o exportar a PDF: horizontal y todo el ancho (tablas + gráficas) en una página."""
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    estilos = SimpleNamespace(hdr_font=hdr_font, hdr_fill=hdr_fill, titulo=titulo, dato=dato, calc=calc,
                              negrita=Font(name=F, bold=True), borde=borde, centro=centro, estilo=estilo)

    wb = Workbook()
    wb.remove(wb.active)
    hojas = {}                                   # tamaño -> (nombre de la hoja, primera y última fila de ejecuciones)
    # columnas: A hilos | B generaciones | C ejecuciones | D..G tiempos V1..V4 | H..J speedups | K más rápido
    for tamano in sorted({f["tamano"] for f in filas}):
        propias = [f for f in filas if f["tamano"] == tamano]
        configs = [k for k in configuraciones(propias)]
        ws = wb.create_sheet(f"{tamano}x{tamano}")

        # --- tabla 1: promedios (arriba)
        ws.cell(row=1, column=1, value=f"Matriz {tamano} × {tamano} — promedios por configuración").font = titulo
        encabezado(ws, 2, ["Hilos/proc.", "Generaciones", "Ejecuciones", "V1 Secuencial", "V2 Hilos",
                           "V3 Procesos", "V4 NumPy", "Speedup V2", "Speedup V3", "Speedup V4", "Más rápido",
                           "Configuración", "Base V1 (1×)"])
        fila_titulo_ej = 3 + len(configs) + 1            # deja una fila en blanco
        fila_head_ej = fila_titulo_ej + 1
        primera = fila_head_ej + 1
        ultima = fila_head_ej + len(propias)

        for r, (_, hilos, gens) in enumerate(configs, 3):
            for c, valor in ((1, hilos), (2, gens)):
                estilo(ws.cell(row=r, column=c, value=valor), dato)
            crit = (f"$A${primera}:$A${ultima},$A{r},$B${primera}:$B${ultima},$B{r}")
            ws.cell(row=r, column=3, value=f"=COUNTIFS({crit})")
            for j, col in enumerate("DEFG"):
                ws.cell(row=r, column=4 + j,
                        value=f"=AVERAGEIFS(${col}${primera}:${col}${ultima},{crit})")
            ws.cell(row=r, column=8, value=f"=$D{r}/E{r}")
            ws.cell(row=r, column=9, value=f"=$D{r}/F{r}")
            ws.cell(row=r, column=10, value=f"=$D{r}/G{r}")
            ws.cell(row=r, column=11, value=f"=INDEX($D$2:$G$2,MATCH(MIN(D{r}:G{r}),D{r}:G{r},0))")
            ws.cell(row=r, column=12, value=f'=A{r}&" hilos · "&B{r}&" gen."')
            ws.cell(row=r, column=13, value=1)           # valor fijo: línea de referencia de la gráfica de speedup
            for c in range(3, 14):
                formato = '0.000" s"' if 4 <= c <= 7 else ('0.00"×"' if c in (8, 9, 10, 13) else None)
                estilo(ws.cell(row=r, column=c), dato if c == 13 else calc, formato)
            if r % 2 == 0:
                for c in range(1, 14):
                    ws.cell(row=r, column=c).fill = banda

        # --- tabla 2: todas las ejecuciones (abajo)
        ws.cell(row=fila_titulo_ej, column=1, value=f"Matriz {tamano} × {tamano} — todas las ejecuciones").font = titulo
        encabezado(ws, fila_head_ej, ["Hilos/proc.", "Generaciones", "Ejecución", "V1 Secuencial", "V2 Hilos",
                                      "V3 Procesos", "V4 NumPy", "Speedup V2", "Speedup V3", "Speedup V4",
                                      "Más rápido", "Correctitud"])
        for r, fila in enumerate(propias, primera):
            valores = [fila["trabajadores"], fila["generaciones"], fila["ejecucion"],
                       *[round(fila["tiempos"][v], 4) for v in VERSIONES]]
            for c, valor in enumerate(valores, 1):
                estilo(ws.cell(row=r, column=c, value=valor), dato, '0.00" s"' if c >= 4 else None)
            ws.cell(row=r, column=8, value=f"=$D{r}/E{r}")
            ws.cell(row=r, column=9, value=f"=$D{r}/F{r}")
            ws.cell(row=r, column=10, value=f"=$D{r}/G{r}")
            ws.cell(row=r, column=11, value=f"=INDEX($D${fila_head_ej}:$G${fila_head_ej},MATCH(MIN(D{r}:G{r}),D{r}:G{r},0))")
            ws.cell(row=r, column=12, value="V2, V3 y V4 iguales a V1" if fila["correcto"] else "¡DIFERENTE A V1!")
            for c in range(8, 13):
                estilo(ws.cell(row=r, column=c), calc, '0.00"×"' if c < 11 else None)
            k = configs.index((tamano, fila["trabajadores"], fila["generaciones"]))
            if k % 2 == 0:                               # franjas alternas por configuración
                for c in range(1, 13):
                    ws.cell(row=r, column=c).fill = banda

        notas = ["Datos medidos en azul; cada fila de ejecuciones es una Carrera completa (las 4 versiones a la vez).",
                 "Speedup = tiempo de V1 Secuencial ÷ tiempo de la versión (mayor que 1× = más rápida que V1).",
                 "Las gráficas de la derecha leen la tabla de promedios (columnas D a J, L y M): si cambian los datos, se actualizan.",
                 "'Base V1 (1×)' es un valor fijo (1) que solo sirve de línea de referencia en la gráfica de speedup."]
        for i, t in enumerate(notas):
            ws.cell(row=ultima + 2 + i, column=1, value=t).font = calc
        for i, ancho in enumerate([12, 14, 12, 15, 12, 13, 12, 12, 12, 12, 16, 24, 14], 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = "A3"
        _graficas_tamano(ws, len(configs))
        pagina_horizontal(ws)
        hojas[tamano] = (ws.title, primera, ultima)

    pagina_horizontal(_hoja_resumen(wb, filas, hojas, estilos))
    wb.save(ruta)


def imprimir_resumen(filas):
    """Tabla de promedios en la consola."""
    print("\n=== PROMEDIOS (segundos; speedup contra V1 entre paréntesis) ===")
    print(f"{'Matriz':>9} {'Hilos':>5} {'Gen.':>5} | " + " | ".join(f"{NOMBRES[v]:^15}" for v in VERSIONES))
    for t, p, g in configuraciones(filas):
        grupo = [f for f in filas if (f["tamano"], f["trabajadores"], f["generaciones"]) == (t, p, g)]
        prom = {v: statistics.mean(f["tiempos"][v] for f in grupo) for v in VERSIONES}
        celdas = [f"{prom[v]:.3f}s ({prom[1] / prom[v]:.2f}x)".center(15) for v in VERSIONES]
        print(f"{f'{t}x{t}':>9} {p:>5} {g:>5} | " + " | ".join(celdas))


# --- 3. principal
def main():
    ap = argparse.ArgumentParser(description="Repite la Carrera de la GUI y genera la tabla de resultados.")
    ap.add_argument("--tamanos", type=lista_enteros, default=[100, 500, 1000], help="lados de la matriz, ej. 100,500,1000")
    ap.add_argument("--trabajadores", type=lista_enteros, default=[4, 8, 16], help="hilos/procesos, ej. 4,8,16")
    ap.add_argument("--generaciones", type=lista_enteros, default=[200, 500], help="ej. 200,500")
    ap.add_argument("--repeticiones", type=int, default=5, help="ejecuciones por configuración")
    ap.add_argument("--semilla", type=int, default=config.SEMILLA)
    ap.add_argument("--salida", default=CARPETA, help="carpeta donde se guardan los archivos")
    ap.add_argument("--solo-excel", action="store_true",
                    help="no mide: rehace el Excel (tablas y gráficas) a partir de ejecuciones.csv")
    a = ap.parse_args()

    os.makedirs(a.salida, exist_ok=True)
    ruta_csv = os.path.join(a.salida, "ejecuciones.csv")

    if a.solo_excel:
        if not os.path.exists(ruta_csv):
            print(f"No existe {ruta_csv}: primero hay que medir (corre el script sin --solo-excel).")
            return
        filas = leer_csv(ruta_csv)
        print(f"Se leyeron {len(filas)} ejecuciones de {ruta_csv} (no se mide nada).")
    else:
        print(f"Matrices {a.tamanos} · hilos/proc {a.trabajadores} · generaciones {a.generaciones} · "
              f"{a.repeticiones} ejecuciones por configuración (semilla {a.semilla})\n")
        filas = []
        try:
            medir(filas, a.tamanos, a.trabajadores, a.generaciones, a.repeticiones, a.semilla, ruta_csv)
        except KeyboardInterrupt:
            print("\nInterrumpido: se guardan las ejecuciones que ya terminaron.")
        if not filas:
            print("No hay ejecuciones que guardar.")
            return
        guardar_csv(filas, ruta_csv)
        print(f"\nCSV guardado en: {ruta_csv}")

    try:
        ruta_xlsx = os.path.join(a.salida, "comparacion_resultados.xlsx")
        guardar_xlsx(filas, ruta_xlsx)
        print(f"Excel guardado en: {ruta_xlsx}")
    except ImportError:
        print("No se creó el .xlsx porque falta openpyxl. Instálalo con:  pip install openpyxl  y vuelve a correr.")
    except PermissionError:
        print(f"No se pudo guardar {ruta_xlsx}: probablemente está abierto en Excel. Ciérralo y corre:\n"
              "    python benchmark_carrera.py --solo-excel\n(los datos ya están a salvo en ejecuciones.csv; no hay que medir otra vez)")
    imprimir_resumen(filas)
    if not all(f["correcto"] for f in filas):
        print("\n¡OJO! Alguna ejecución dio una matriz distinta a V1. Revisa la columna Correctitud.")


if __name__ == "__main__":      # obligatorio en Windows: los procesos hijos (spawn) reimportan este archivo
    main()